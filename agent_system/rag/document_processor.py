"""
文档处理模块

处理各种格式的文档，提取文本并进行分块
"""

import os
from typing import List, Dict, Any, Optional
from pathlib import Path
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
import docx

from ..config.settings import RAG_CHUNK_SIZE, RAG_CHUNK_OVERLAP


class DocumentProcessor:
    """
    文档处理类
    
    支持多种文档格式的文本提取和分块处理
    """
    
    def __init__(
        self,
        chunk_size: int = RAG_CHUNK_SIZE,
        chunk_overlap: int = RAG_CHUNK_OVERLAP,
        vl_tools=None
    ):
        """
        初始化文档处理器
        
        Args:
            chunk_size: 文本分块大小
            chunk_overlap: 文本分块重叠大小
            vl_tools: Qwen2.5-VL 工具实例（用于 OCR）
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.vl_tools = vl_tools
        
        # 初始化文本分割器
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", "。", "！", "？", "；", "，", " ", ""]
        )
    
    def process_file(
        self,
        file_path: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> List[Document]:
        """
        处理单个文件
        
        Args:
            file_path: 文件路径
            metadata: 额外的元数据
            
        Returns:
            文档分块列表
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        # 根据文件扩展名选择处理方法
        suffix = file_path.suffix.lower()
        
        if suffix == '.pdf':
            text = self._extract_pdf_text(file_path)
        elif suffix == '.txt':
            text = self._extract_txt_text(file_path)
        elif suffix in ['.docx', '.doc']:
            text = self._extract_docx_text(file_path)
        elif suffix == '.md':
            text = self._extract_txt_text(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {suffix}")
        
        # 准备元数据
        doc_metadata = {
            "source": str(file_path),
            "file_name": file_path.name,
            "file_type": suffix[1:]
        }
        if metadata:
            doc_metadata.update(metadata)
        
        # 分块处理
        chunks = self.text_splitter.split_text(text)
        
        # 创建 Document 对象
        documents = []
        for i, chunk in enumerate(chunks):
            chunk_metadata = doc_metadata.copy()
            chunk_metadata["chunk_id"] = i
            chunk_metadata["total_chunks"] = len(chunks)
            
            documents.append(Document(
                page_content=chunk,
                metadata=chunk_metadata
            ))
        
        return documents
    
    def process_directory(
        self,
        directory: str,
        recursive: bool = True,
        file_extensions: Optional[List[str]] = None
    ) -> List[Document]:
        """
        批量处理目录中的文件
        
        Args:
            directory: 目录路径
            recursive: 是否递归处理子目录
            file_extensions: 要处理的文件扩展名列表（如 ['.pdf', '.txt']）
            
        Returns:
            所有文档的分块列表
        """
        directory = Path(directory)
        
        if not directory.exists():
            raise FileNotFoundError(f"目录不存在: {directory}")
        
        # 默认支持的文件扩展名
        if file_extensions is None:
            file_extensions = ['.pdf', '.txt', '.docx', '.doc', '.md']
        
        # 收集所有文件
        files = []
        if recursive:
            for ext in file_extensions:
                files.extend(directory.rglob(f"*{ext}"))
        else:
            for ext in file_extensions:
                files.extend(directory.glob(f"*{ext}"))
        
        print(f"找到 {len(files)} 个文件待处理")
        
        # 处理所有文件
        all_documents = []
        for i, file_path in enumerate(files, 1):
            try:
                print(f"[{i}/{len(files)}] 处理文件: {file_path.name}")
                documents = self.process_file(file_path)
                all_documents.extend(documents)
                print(f"  ✓ 生成 {len(documents)} 个文档块")
            except Exception as e:
                print(f"  ✗ 处理失败: {str(e)}")
        
        print(f"\n总计生成 {len(all_documents)} 个文档块")
        return all_documents
    
    def _extract_pdf_text(self, file_path: Path) -> str:
        """
        从 PDF 提取文本
        
        优先使用 Qwen2.5-VL 的 OCR，如果不可用则使用 pypdf
        
        Args:
            file_path: PDF 文件路径
            
        Returns:
            提取的文本
        """
        # 尝试使用 Qwen2.5-VL OCR
        if self.vl_tools:
            print("使用 Qwen2.5-VL OCR")
            try:
                pages = self.vl_tools.extract_text(file_path=str(file_path))
                if pages and not any('error' in page for page in pages):
                    # 合并所有页面的文本
                    text = "\n\n".join([
                        page.get('text', '') for page in pages
                        if 'text' in page
                    ])
                    if text.strip():
                        return text
            except Exception as e:
                print(f"  Qwen2.5-VL OCR 失败，使用 pypdf: {str(e)}")
        
        # 使用 pypdf 提取文本
        try:
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n\n"
            return text.strip()
        except Exception as e:
            raise Exception(f"PDF 文本提取失败: {str(e)}")
    
    def _extract_txt_text(self, file_path: Path) -> str:
        """
        从文本文件提取文本
        
        Args:
            file_path: 文本文件路径
            
        Returns:
            文件内容
        """
        try:
            # 尝试多种编码
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        return f.read()
                except UnicodeDecodeError:
                    continue
            
            raise Exception("无法识别文件编码")
        except Exception as e:
            raise Exception(f"文本文件读取失败: {str(e)}")
    
    def _extract_docx_text(self, file_path: Path) -> str:
        """
        从 Word 文档提取文本
        
        Args:
            file_path: Word 文档路径
            
        Returns:
            提取的文本
        """
        try:
            doc = docx.Document(file_path)
            text = "\n\n".join([paragraph.text for paragraph in doc.paragraphs])
            return text.strip()
        except Exception as e:
            raise Exception(f"Word 文档读取失败: {str(e)}")


