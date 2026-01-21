"""
嵌入流水线模块

将分块器、元数据提取器、嵌入服务组合成完整的处理流水线
支持灵活切换各个组件进行效果对比
"""

import time
from pathlib import Path
from typing import List, Dict, Any, Optional, Union, Type
from dataclasses import dataclass, field

from langchain_core.documents import Document

from ..chunkers.base import ChunkerBase
from ..chunkers.legal_chunker import LegalChunker
from ..chunkers.recursive_chunker import RecursiveChunker
from ..extractors.base import ExtractorBase
from ..extractors.legal_metadata import LegalMetadataExtractor
from .service import EmbeddingService, create_embedding_service
from .cache import EmbeddingCache
from .batcher import BatchProcessor, BatchResult


@dataclass
class PipelineResult:
    """流水线处理结果"""
    documents: List[Document]
    embeddings: List[List[float]]
    metadata: Dict[str, Any]
    stats: Dict[str, Any] = field(default_factory=dict)


@dataclass
class PipelineConfig:
    """流水线配置"""
    # 分块器配置
    chunker_type: str = "legal"  # "legal", "recursive", "semantic"
    chunker_params: Dict[str, Any] = field(default_factory=dict)
    
    # 元数据提取器配置
    extractor_type: str = "legal"
    extractor_params: Dict[str, Any] = field(default_factory=dict)
    
    # 嵌入服务配置
    embedding_provider: str = "qwen"  # "qwen", "openai", "local"
    embedding_model: Optional[str] = None
    embedding_params: Dict[str, Any] = field(default_factory=dict)
    
    # 缓存配置
    enable_cache: bool = True
    cache_dir: Optional[str] = "./data/embedding_cache"
    
    # 批处理配置
    batch_size: int = 100
    max_workers: int = 4


class EmbeddingPipeline:
    """
    嵌入处理流水线
    
    将文档处理的各个阶段组合成完整流水线：
    1. 读取文档
    2. 提取元数据
    3. 文档分块
    4. 向量嵌入
    
    支持灵活切换各个组件：
    - 分块器：LegalChunker, RecursiveChunker, SemanticChunker
    - 提取器：LegalMetadataExtractor
    - 嵌入服务：Qwen, OpenAI, 本地模型
    
    使用示例:
        # 使用默认配置
        pipeline = EmbeddingPipeline()
        result = pipeline.process_file("法律文档.docx")
        
        # 自定义配置
        config = PipelineConfig(
            chunker_type="recursive",
            embedding_provider="openai"
        )
        pipeline = EmbeddingPipeline(config=config)
        
        # 切换分块器进行对比
        pipeline.set_chunker("semantic")
        result2 = pipeline.process_file("法律文档.docx")
    """
    
    # 分块器注册表
    CHUNKER_REGISTRY: Dict[str, Type[ChunkerBase]] = {
        "legal": LegalChunker,
        "recursive": RecursiveChunker,
    }
    
    # 提取器注册表
    EXTRACTOR_REGISTRY: Dict[str, Type[ExtractorBase]] = {
        "legal": LegalMetadataExtractor,
    }
    
    def __init__(
        self,
        config: Optional[PipelineConfig] = None,
        chunker: Optional[ChunkerBase] = None,
        extractor: Optional[ExtractorBase] = None,
        embedding_service: Optional[EmbeddingService] = None,
        cache: Optional[EmbeddingCache] = None
    ):
        """
        初始化流水线
        
        Args:
            config: 流水线配置，如果为 None 则使用默认配置
            chunker: 自定义分块器实例
            extractor: 自定义元数据提取器实例
            embedding_service: 自定义嵌入服务实例
            cache: 自定义缓存实例
        """
        self.config = config or PipelineConfig()
        
        # 初始化组件
        self._chunker = chunker
        self._extractor = extractor
        self._embedding_service = embedding_service
        self._cache = cache
        self._batch_processor = None
        
        # 延迟初始化标志
        self._initialized = False
    
    def _ensure_initialized(self):
        """确保所有组件已初始化"""
        if self._initialized:
            return
        
        # 初始化分块器
        if self._chunker is None:
            self._chunker = self._create_chunker(
                self.config.chunker_type,
                self.config.chunker_params
            )
        
        # 初始化提取器
        if self._extractor is None:
            self._extractor = self._create_extractor(
                self.config.extractor_type,
                self.config.extractor_params
            )
        
        # 初始化嵌入服务
        if self._embedding_service is None:
            self._embedding_service = create_embedding_service(
                provider=self.config.embedding_provider,
                model=self.config.embedding_model,
                **self.config.embedding_params
            )
        
        # 初始化缓存
        if self._cache is None and self.config.enable_cache:
            self._cache = EmbeddingCache(
                cache_dir=self.config.cache_dir
            )
        
        # 初始化批处理器
        self._batch_processor = BatchProcessor(
            embedding_service=self._embedding_service,
            batch_size=self.config.batch_size,
            max_workers=self.config.max_workers
        )
        
        self._initialized = True
    
    def _create_chunker(
        self,
        chunker_type: str,
        params: Dict[str, Any]
    ) -> ChunkerBase:
        """创建分块器实例"""
        # 尝试动态加载 SemanticChunker
        if chunker_type == "semantic":
            try:
                from ..chunkers.semantic_chunker import SemanticChunker
                return SemanticChunker(**params)
            except ImportError as e:
                print(f"无法加载 SemanticChunker: {e}")
                print("回退到 LegalChunker")
                chunker_type = "legal"
        
        if chunker_type not in self.CHUNKER_REGISTRY:
            raise ValueError(f"未知的分块器类型: {chunker_type}")
        
        return self.CHUNKER_REGISTRY[chunker_type](**params)
    
    def _create_extractor(
        self,
        extractor_type: str,
        params: Dict[str, Any]
    ) -> ExtractorBase:
        """创建提取器实例"""
        if extractor_type not in self.EXTRACTOR_REGISTRY:
            raise ValueError(f"未知的提取器类型: {extractor_type}")
        
        return self.EXTRACTOR_REGISTRY[extractor_type](**params)
    
    def set_chunker(
        self,
        chunker_type: str,
        **params
    ) -> "EmbeddingPipeline":
        """
        切换分块器
        
        Args:
            chunker_type: 分块器类型
            **params: 分块器参数
            
        Returns:
            self，支持链式调用
        """
        self._chunker = self._create_chunker(chunker_type, params)
        return self
    
    def set_extractor(
        self,
        extractor_type: str,
        **params
    ) -> "EmbeddingPipeline":
        """
        切换提取器
        
        Args:
            extractor_type: 提取器类型
            **params: 提取器参数
            
        Returns:
            self，支持链式调用
        """
        self._extractor = self._create_extractor(extractor_type, params)
        return self
    
    def set_embedding_service(
        self,
        provider: str,
        **params
    ) -> "EmbeddingPipeline":
        """
        切换嵌入服务
        
        Args:
            provider: 嵌入服务提供商
            **params: 嵌入服务参数
            
        Returns:
            self，支持链式调用
        """
        self._embedding_service = create_embedding_service(provider, **params)
        
        # 重新创建批处理器
        if self._batch_processor is not None:
            self._batch_processor = BatchProcessor(
                embedding_service=self._embedding_service,
                batch_size=self.config.batch_size,
                max_workers=self.config.max_workers
            )
        
        return self
    
    @property
    def chunker(self) -> ChunkerBase:
        """获取当前分块器"""
        self._ensure_initialized()
        return self._chunker
    
    @property
    def extractor(self) -> ExtractorBase:
        """获取当前提取器"""
        self._ensure_initialized()
        return self._extractor
    
    @property
    def embedding_service(self) -> EmbeddingService:
        """获取当前嵌入服务"""
        self._ensure_initialized()
        return self._embedding_service
    
    def process_text(
        self,
        text: str,
        metadata: Optional[Dict[str, Any]] = None,
        source: str = "unknown"
    ) -> PipelineResult:
        """
        处理文本
        
        Args:
            text: 文本内容
            metadata: 可选的初始元数据
            source: 来源标识
            
        Returns:
            PipelineResult 对象
        """
        self._ensure_initialized()
        
        start_time = time.time()
        stats = {}
        
        # 1. 分块
        chunk_start = time.time()
        base_metadata = metadata or {}
        base_metadata["source"] = source
        
        documents = self._chunker.chunk(text, base_metadata)
        stats["chunk_time"] = time.time() - chunk_start
        stats["chunk_count"] = len(documents)
        
        # 2. 嵌入
        embed_start = time.time()
        batch_result = self._batch_processor.process(documents, self._cache)
        stats["embed_time"] = time.time() - embed_start
        stats["embed_success"] = batch_result.success_count
        stats["embed_errors"] = batch_result.error_count
        
        # 3. 汇总统计
        stats["total_time"] = time.time() - start_time
        
        if self._cache:
            cache_stats = self._cache.get_stats()
            stats["cache_hit_rate"] = cache_stats["hit_rate"]
        
        return PipelineResult(
            documents=documents,
            embeddings=batch_result.embeddings,
            metadata=base_metadata,
            stats=stats
        )
    
    def process_file(
        self,
        file_path: Union[str, Path],
        extract_metadata: bool = True
    ) -> PipelineResult:
        """
        处理文件
        
        Args:
            file_path: 文件路径
            extract_metadata: 是否提取元数据
            
        Returns:
            PipelineResult 对象
        """
        self._ensure_initialized()
        
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        # 读取文件内容
        text = self._read_file(file_path)
        
        # 提取元数据
        metadata = {}
        if extract_metadata:
            metadata = self._extractor.extract(str(file_path), text)
        
        metadata["source"] = str(file_path)
        metadata["file_name"] = file_path.name
        
        # 处理
        return self.process_text(text, metadata, str(file_path))
    
    def process_directory(
        self,
        directory: Union[str, Path],
        recursive: bool = True,
        file_extensions: Optional[List[str]] = None,
        progress_callback: Optional[callable] = None
    ) -> List[PipelineResult]:
        """
        处理目录下的所有文件
        
        Args:
            directory: 目录路径
            recursive: 是否递归处理子目录
            file_extensions: 要处理的文件扩展名列表
            progress_callback: 进度回调 (当前索引, 总数, 文件名)
            
        Returns:
            PipelineResult 列表
        """
        self._ensure_initialized()
        
        directory = Path(directory)
        
        if not directory.exists():
            raise FileNotFoundError(f"目录不存在: {directory}")
        
        # 默认支持的扩展名
        if file_extensions is None:
            file_extensions = [".pdf", ".docx", ".doc", ".txt", ".md"]
        
        # 收集文件
        files = []
        if recursive:
            for ext in file_extensions:
                files.extend(directory.rglob(f"*{ext}"))
        else:
            for ext in file_extensions:
                files.extend(directory.glob(f"*{ext}"))
        
        # 处理文件
        results = []
        for i, file_path in enumerate(files):
            if progress_callback:
                progress_callback(i, len(files), file_path.name)
            
            try:
                result = self.process_file(file_path)
                results.append(result)
            except Exception as e:
                print(f"处理文件失败 {file_path.name}: {e}")
        
        return results
    
    def _read_file(self, file_path: Path) -> str:
        """读取文件内容"""
        suffix = file_path.suffix.lower()
        
        if suffix == ".docx":
            import docx
            doc = docx.Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs])
        
        elif suffix == ".txt" or suffix == ".md":
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        return f.read()
                except UnicodeDecodeError:
                    continue
            raise Exception("无法识别文件编码")
        
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n\n"
            return text.strip()
        
        elif suffix == ".doc":
            # 尝试使用 docx 读取（可能失败）
            try:
                import docx
                doc = docx.Document(file_path)
                return "\n".join([p.text for p in doc.paragraphs])
            except Exception:
                raise Exception("不支持的 .doc 格式，请转换为 .docx")
        
        else:
            raise ValueError(f"不支持的文件格式: {suffix}")
    
    def compare_chunkers(
        self,
        text: str,
        chunker_types: List[str] = None,
        metadata: Optional[Dict[str, Any]] = None
    ) -> Dict[str, PipelineResult]:
        """
        对比不同分块器的效果
        
        Args:
            text: 文本内容
            chunker_types: 要对比的分块器类型列表
            metadata: 初始元数据
            
        Returns:
            {分块器类型: PipelineResult} 字典
        """
        self._ensure_initialized()
        
        if chunker_types is None:
            chunker_types = ["legal", "recursive"]
        
        results = {}
        original_chunker = self._chunker
        
        for chunker_type in chunker_types:
            try:
                self.set_chunker(chunker_type)
                result = self.process_text(text, metadata, f"compare_{chunker_type}")
                results[chunker_type] = result
            except Exception as e:
                print(f"分块器 {chunker_type} 失败: {e}")
        
        # 恢复原始分块器
        self._chunker = original_chunker
        
        return results
    
    def get_stats(self) -> Dict[str, Any]:
        """获取流水线统计信息"""
        stats = {
            "chunker_type": type(self._chunker).__name__ if self._chunker else None,
            "extractor_type": type(self._extractor).__name__ if self._extractor else None,
            "embedding_provider": self.config.embedding_provider,
            "embedding_model": self.config.embedding_model,
        }
        
        if self._cache:
            stats["cache"] = self._cache.get_stats()
        
        return stats
