"""
预处理层测试

使用 data/law-data 目录下的真实法律文件进行测试
展示详细的分块结果和内容
"""

import os
import sys
from pathlib import Path

# 添加项目根目录到 Python 路径
project_root = Path(__file__).parent.parent.parent
sys.path.insert(0, str(project_root))

# 直接导入分块器和提取器，避免加载整个 rag 模块（会触发 dashscope 导入）
from agent_system.rag.chunkers.legal_chunker import LegalChunker
from agent_system.rag.chunkers.recursive_chunker import RecursiveChunker
from agent_system.rag.extractors.legal_metadata import LegalMetadataExtractor


def read_docx(file_path: Path) -> str:
    """读取 docx 文件内容"""
    import docx
    doc = docx.Document(file_path)
    return "\n".join([p.text for p in doc.paragraphs])


def read_doc(file_path: Path) -> str:
    """读取 doc 文件内容（尝试使用 docx 库，可能失败）"""
    try:
        import docx
        doc = docx.Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs])
    except Exception:
        # 对于旧版 .doc 文件，返回空（需要其他库如 antiword）
        return ""


def indent_text(text: str, indent: str = "    ") -> str:
    """为多行文本添加缩进"""
    lines = text.split("\n")
    return "\n".join(indent + line for line in lines)


def test_single_file(file_path: Path, extractor: LegalMetadataExtractor, chunker: LegalChunker):
    """测试单个文件的元数据提取和分块"""
    print("\n" + "=" * 80)
    print(f"📄 文件: {file_path.name}")
    print("=" * 80)
    
    # 读取文件内容
    try:
        if file_path.suffix.lower() == ".docx":
            text = read_docx(file_path)
        else:
            text = read_doc(file_path)
        
        if not text:
            print("  ⚠️ 无法读取文件内容")
            return False
            
    except Exception as e:
        print(f"  ❌ 读取失败: {e}")
        return False
    
    print(f"原文长度: {len(text)} 字符")
    
    # 1. 提取元数据
    print("\n" + "-" * 40)
    print("【元数据提取结果】")
    print("-" * 40)
    
    metadata = extractor.extract(str(file_path), text)
    
    print(f"  法规名称: {metadata['law_name']}")
    print(f"  法规类型: {metadata['law_type']} ({metadata['law_type_en']})")
    print(f"  地区: {metadata['region']}")
    print(f"  地区级别: {metadata['region_level']} ({metadata['region_level_en']})")
    print(f"  主题分类: {metadata['topics']}")
    print(f"  颁布日期: {metadata['publish_date'] or '未提取到'}")
    print(f"  生效日期: {metadata['effective_date'] or '未提取到'}")
    print(f"  颁布机关: {metadata['issuing_authority'] or '未提取到'}")
    
    # 2. 智能分块
    print("\n" + "-" * 40)
    print("【智能分块结果】")
    print("-" * 40)
    
    chunks = chunker.chunk(text, metadata)
    
    print(f"  分块数量: {len(chunks)} 块")
    
    # 统计章节分布
    chapters = {}
    for chunk in chunks:
        ch = chunk.metadata.get("chapter", "") or "(无章节/序言)"
        chapters[ch] = chapters.get(ch, 0) + 1
    
    print(f"\n  章节分布:")
    for ch, count in chapters.items():
        print(f"    • {ch}: {count} 块")
    
    # 显示所有分块的详细内容
    print("\n" + "-" * 40)
    print("【分块详细内容】")
    print("-" * 40)
    
    for i, chunk in enumerate(chunks):
        print(f"\n┌{'─' * 76}┐")
        print(f"│ 分块 {i + 1}/{len(chunks)}")
        print(f"├{'─' * 76}┤")
        print(f"│ 章: {chunk.metadata.get('chapter', 'N/A')}")
        print(f"│ 节: {chunk.metadata.get('section', 'N/A')}")
        print(f"│ 条款号: {chunk.metadata.get('article_num', 'N/A')}")
        print(f"│ 上下文: {chunk.metadata.get('parent_context', 'N/A')}")
        print(f"│ 内容长度: {len(chunk.page_content)} 字符")
        print(f"├{'─' * 76}┤")
        
        # 显示完整内容（限制最大长度）
        content = chunk.page_content
        if len(content) > 800:
            content = content[:800] + "\n... (内容过长，已截断) ..."
        
        # 格式化内容显示
        for line in content.split("\n"):
            if line.strip():
                # 每行最多显示74个字符
                while len(line) > 74:
                    print(f"│ {line[:74]}")
                    line = line[74:]
                print(f"│ {line}")
        
        print(f"└{'─' * 76}┘")
    
    return True


def test_with_law_data():
    """使用 law-data 目录下的真实法律文件进行测试"""
    print("\n" + "=" * 80)
    print("       法律文档预处理测试（使用 law-data 真实文件）")
    print("=" * 80)
    
    # 定位 law-data 目录
    law_data_dir = project_root / "data" / "law-data"
    
    if not law_data_dir.exists():
        print(f"❌ 测试目录不存在: {law_data_dir}")
        return False
    
    # 初始化组件
    print("\n初始化组件...")
    extractor = LegalMetadataExtractor()
    chunker = LegalChunker(max_chunk_size=1500, min_chunk_size=100)
    
    # 选择几个代表性文件进行测试
    test_files = [
        "深圳经济特区消防条例_.docx",
        "海南经济特区水条例_.docx",
        "海北藏族自治州义务教育条例_.docx",
    ]
    
    # 如果指定文件不存在，则使用目录中的前3个文件
    available_files = []
    for f in test_files:
        fp = law_data_dir / f
        if fp.exists():
            available_files.append(fp)
    
    if len(available_files) < 3:
        # 补充其他 docx 文件
        for fp in law_data_dir.glob("*.docx"):
            if fp not in available_files:
                available_files.append(fp)
            if len(available_files) >= 3:
                break
    
    if not available_files:
        print("❌ 没有找到可用的法律文件")
        return False
    
    print(f"找到 {len(available_files)} 个测试文件")
    
    success_count = 0
    
    for file_path in available_files:
        if test_single_file(file_path, extractor, chunker):
            success_count += 1
    
    print("\n" + "=" * 80)
    print(f"测试完成: {success_count}/{len(available_files)} 个文件处理成功")
    print("=" * 80)
    
    return success_count > 0


def test_chunker_comparison():
    """对比法律分块器和通用分块器的效果"""
    print("\n" + "=" * 80)
    print("       分块器效果对比")
    print("=" * 80)
    
    law_data_dir = project_root / "data" / "law-data"
    
    # 找一个法律文件
    test_file = None
    for fp in law_data_dir.glob("*.docx"):
        test_file = fp
        break
    
    if not test_file:
        print("❌ 没有找到测试文件")
        return False
    
    text = read_docx(test_file)
    if not text:
        print("❌ 无法读取文件")
        return False
    
    print(f"\n测试文件: {test_file.name}")
    print(f"原文长度: {len(text)} 字符")
    
    # 法律分块器
    legal_chunker = LegalChunker(max_chunk_size=1000, min_chunk_size=100)
    legal_chunks = legal_chunker.chunk(text, {"source": str(test_file)})
    
    # 通用分块器
    recursive_chunker = RecursiveChunker(chunk_size=800, chunk_overlap=100)
    recursive_chunks = recursive_chunker.chunk(text, {"source": str(test_file)})
    
    print(f"\n【法律分块器 (LegalChunker)】")
    print(f"  分块数: {len(legal_chunks)}")
    if legal_chunks:
        print(f"  平均长度: {sum(len(c.page_content) for c in legal_chunks) // len(legal_chunks)} 字符")
    
    # 检查条款完整性
    complete_articles = sum(1 for c in legal_chunks if c.metadata.get("article_num"))
    print(f"  包含完整条款号的块: {complete_articles}")
    
    print(f"\n【通用分块器 (RecursiveChunker)】")
    print(f"  分块数: {len(recursive_chunks)}")
    if recursive_chunks:
        print(f"  平均长度: {sum(len(c.page_content) for c in recursive_chunks) // len(recursive_chunks)} 字符")
    
    print("\n✓ 对比完成")
    return True


def test_metadata_extraction_batch():
    """批量测试元数据提取"""
    print("\n" + "=" * 80)
    print("       批量元数据提取统计")
    print("=" * 80)
    
    law_data_dir = project_root / "data" / "law-data"
    extractor = LegalMetadataExtractor()
    
    # 统计
    stats = {
        "total": 0,
        "with_region": 0,
        "with_type": 0,
        "with_topics": 0,
        "with_date": 0,
        "regions": {},
        "types": {},
        "region_levels": {},
    }
    
    # 处理所有 docx 文件（最多20个）
    files = list(law_data_dir.glob("*.docx"))[:20]
    
    print(f"\n处理 {len(files)} 个文件...")
    
    for fp in files:
        try:
            text = read_docx(fp)
            if not text:
                continue
            
            metadata = extractor.extract(str(fp), text)
            
            stats["total"] += 1
            
            if metadata["region"]:
                stats["with_region"] += 1
                stats["regions"][metadata["region"]] = stats["regions"].get(metadata["region"], 0) + 1
            
            if metadata["law_type"]:
                stats["with_type"] += 1
                stats["types"][metadata["law_type"]] = stats["types"].get(metadata["law_type"], 0) + 1
            
            if metadata["topics"]:
                stats["with_topics"] += 1
            
            if metadata["publish_date"] or metadata["effective_date"]:
                stats["with_date"] += 1
            
            if metadata["region_level"]:
                stats["region_levels"][metadata["region_level"]] = \
                    stats["region_levels"].get(metadata["region_level"], 0) + 1
                
        except Exception as e:
            print(f"  ⚠️ {fp.name}: {e}")
    
    print(f"\n【统计结果】")
    print(f"  处理文件数: {stats['total']}")
    print(f"  提取到地区: {stats['with_region']} ({stats['with_region']*100//max(stats['total'],1)}%)")
    print(f"  提取到类型: {stats['with_type']} ({stats['with_type']*100//max(stats['total'],1)}%)")
    print(f"  提取到主题: {stats['with_topics']} ({stats['with_topics']*100//max(stats['total'],1)}%)")
    print(f"  提取到日期: {stats['with_date']} ({stats['with_date']*100//max(stats['total'],1)}%)")
    
    print(f"\n【地区分布】")
    for region, count in sorted(stats["regions"].items(), key=lambda x: -x[1])[:10]:
        print(f"    {region}: {count}")
    
    print(f"\n【地区级别分布】")
    for level, count in sorted(stats["region_levels"].items(), key=lambda x: -x[1]):
        print(f"    {level}: {count}")
    
    print(f"\n【类型分布】")
    for law_type, count in sorted(stats["types"].items(), key=lambda x: -x[1]):
        print(f"    {law_type}: {count}")
    
    print("\n✓ 统计完成")
    return True


def run_all_tests():
    """运行所有测试"""
    print("\n" + "=" * 80)
    print("       预处理层完整测试（使用 law-data 真实文件）")
    print("=" * 80)
    
    tests = [
        ("1. 真实法律文件详细测试", test_with_law_data),
        ("2. 分块器效果对比", test_chunker_comparison),
        ("3. 批量元数据提取统计", test_metadata_extraction_batch),
    ]
    
    passed = 0
    failed = 0
    
    for name, test_func in tests:
        print(f"\n\n{'#' * 80}")
        print(f"# {name}")
        print(f"{'#' * 80}")
        
        try:
            result = test_func()
            if result:
                passed += 1
                print(f"\n✓ {name} - 通过")
            else:
                failed += 1
                print(f"\n✗ {name} - 失败")
        except Exception as e:
            failed += 1
            print(f"\n✗ {name} - 异常")
            print(f"  错误: {e}")
            import traceback
            traceback.print_exc()
    
    print("\n" + "=" * 80)
    print(f"总测试结果: {passed} 通过, {failed} 失败")
    print("=" * 80)
    
    return failed == 0


if __name__ == "__main__":
    success = run_all_tests()
    sys.exit(0 if success else 1)
